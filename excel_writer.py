from __future__ import annotations

from collections import defaultdict
from pathlib import Path

import xlsxwriter
from openpyxl import load_workbook

from models import ImageData, MatchedPart


FORMULA_ERRORS = ("#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A")


def _result_label(difference_abs: float) -> str:
    if difference_abs <= 0.01:
        return "一致"
    if difference_abs <= 1:
        return "差异≤1KG"
    return "有差异"


def write_result(path: Path, image: ImageData, matches: tuple[MatchedPart, ...]) -> None:
    if not matches:
        raise ValueError("没有可写入的匹配零件")
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = xlsxwriter.Workbook(path)
    workbook.set_properties({"title": f"{image.main_name}_完成 拆图重量明细"})
    workbook.set_custom_property("原文件名", image.original_filename)
    worksheet = workbook.add_worksheet("拆图明细")
    worksheet.hide_gridlines(2)
    worksheet.set_column("A:A", 8)
    worksheet.set_column("B:B", 24)
    worksheet.set_column("C:C", 26)
    worksheet.set_column("D:D", 10)
    worksheet.set_column("E:E", 12)
    worksheet.set_column("F:F", 14)
    worksheet.set_column("G:G", 17)
    worksheet.set_column("H:H", 16)
    worksheet.set_column("I:I", 17)
    worksheet.set_column("J:J", 3)
    worksheet.set_column("K:K", 8)
    worksheet.set_column("L:L", 24)
    worksheet.set_column("M:M", 17)

    base_font = {"font_name": "Carlito", "font_size": 11, "font_color": "#1F2937"}
    title = workbook.add_format(
        {**base_font, "font_size": 16, "bold": True, "font_color": "#FFFFFF", "bg_color": "#155667", "align": "center", "valign": "vcenter"}
    )
    meta = workbook.add_format(
        {**base_font, "bold": True, "bg_color": "#E5F0F3", "align": "center", "valign": "vcenter"}
    )
    header = workbook.add_format(
        {**base_font, "bold": True, "font_color": "#FFFFFF", "bg_color": "#22798A", "align": "center", "valign": "vcenter", "text_wrap": True, "border": 1}
    )
    row_text = workbook.add_format(
        {**base_font, "bg_color": "#F4F8F9", "align": "center", "valign": "vcenter", "border": 1}
    )
    row_integer = workbook.add_format(
        {**base_font, "bg_color": "#F4F8F9", "align": "center", "valign": "vcenter", "border": 1, "num_format": "0"}
    )
    row_tonnes = workbook.add_format(
        {**base_font, "bg_color": "#F4F8F9", "align": "right", "valign": "vcenter", "border": 1, "num_format": "0.000000"}
    )
    row_kg = workbook.add_format(
        {**base_font, "bg_color": "#F4F8F9", "align": "right", "valign": "vcenter", "border": 1, "num_format": "#,##0.00"}
    )
    summary_title = workbook.add_format(
        {**base_font, "font_size": 12, "bold": True, "font_color": "#FFFFFF", "bg_color": "#155667", "align": "center", "valign": "vcenter"}
    )
    summary_label = workbook.add_format(
        {**base_font, "bold": True, "font_color": "#374151", "bg_color": "#E6F0F2", "border": 1}
    )
    summary_value = workbook.add_format({**base_font, "align": "right", "border": 1, "num_format": "#,##0.00"})
    summary_text = workbook.add_format({**base_font, "align": "center", "border": 1})
    note = workbook.add_format(
        {**base_font, "font_size": 10, "font_color": "#9A5B00", "bg_color": "#FFF7E6", "valign": "vcenter", "text_wrap": True}
    )
    source_note = workbook.add_format(
        {**base_font, "font_size": 9, "font_color": "#4B5563", "bg_color": "#F3F4F6", "text_wrap": True}
    )

    worksheet.set_row(0, 30)
    worksheet.merge_range("A1:I1", f"{image.main_name}_完成 拆图重量明细", title)
    worksheet.set_row(1, 24)
    metadata = ["图片主名称", image.main_name, "程序号", image.program_no, "零件条数", len(image.parts), "核验结果", "通过"]
    for column, value in enumerate(metadata):
        worksheet.write(1, column, value, meta)

    headers = ["序号", "订单号", "图号", "厚度", "基础件数", "图片拆分数量", "基础表总重量（T）", "单件重量（T）", "本次拆分重量（KG）"]
    worksheet.set_row(3, 28)
    worksheet.write_row(3, 0, headers, header)
    for offset, match in enumerate(matches):
        row = 4 + offset
        excel_row = row + 1
        worksheet.set_row(row, 22)
        worksheet.write_number(row, 0, match.image.index, row_integer)
        worksheet.write(row, 1, match.image.order_no, row_text)
        worksheet.write(row, 2, match.image.drawing_no, row_text)
        worksheet.write_number(row, 3, match.image.thickness, row_integer)
        worksheet.write_number(row, 4, match.image.base_quantity, row_integer)
        worksheet.write_number(row, 5, match.image.split_quantity, row_integer)
        worksheet.write_number(row, 6, match.total_weight_t, row_tonnes)
        worksheet.write_formula(row, 7, f"=G{excel_row}/E{excel_row}", row_tonnes, match.unit_weight_t)
        worksheet.write_formula(row, 8, f"=H{excel_row}*F{excel_row}*1000", row_kg, match.split_weight_kg)

    first_data_row = 5
    last_data_row = 4 + len(matches)
    totals: dict[str, float] = defaultdict(float)
    for match in matches:
        totals[match.image.order_no] += match.split_weight_kg
    worksheet.merge_range("K1:M1", "当前图片订单重量汇总（KG）", summary_title)
    worksheet.write_row(1, 10, ["序号", "订单号", "本次拆分重量（KG）"], header)
    for offset, (order_no, total) in enumerate(totals.items()):
        row = 2 + offset
        excel_row = row + 1
        worksheet.write_number(row, 10, offset + 1, row_integer)
        worksheet.write(row, 11, order_no, row_text)
        formula = f'=SUMIF($B${first_data_row}:$B${last_data_row},L{excel_row},$I${first_data_row}:$I${last_data_row})'
        worksheet.write_formula(row, 12, formula, summary_value, total)
    total_row = 2 + len(totals)
    worksheet.merge_range(total_row, 10, total_row, 11, "当前图片合计", summary_label)
    split_total = sum(match.split_weight_kg for match in matches)
    worksheet.write_formula(total_row, 12, f"=SUM(M3:M{total_row})", summary_value, split_total)

    summary_row = max(7, total_row + 2)
    difference = split_total - image.marked_weight_kg
    absolute = abs(difference)
    summary_items = [
        ("图片标注重量（KG）", image.marked_weight_kg, None),
        ("拆分重量合计（KG）", split_total, f"=M{total_row + 1}"),
        ("差额：拆分－图片（KG）", difference, f'=IF(M{summary_row + 1}="","",M{summary_row + 2}-M{summary_row + 1})'),
        ("差额绝对值（KG）", absolute, f'=IF(M{summary_row + 3}="","",ABS(M{summary_row + 3}))'),
    ]
    for offset, (label, value, formula) in enumerate(summary_items):
        row = summary_row + offset
        worksheet.merge_range(row, 10, row, 11, label, summary_label)
        if formula:
            worksheet.write_formula(row, 12, formula, summary_value, value)
        else:
            worksheet.write_number(row, 12, value, summary_value)
    conclusion_row = summary_row + 4
    conclusion = _result_label(absolute)
    worksheet.merge_range(conclusion_row, 10, conclusion_row, 11, "复核结论", summary_label)
    worksheet.write_formula(
        conclusion_row,
        12,
        f'=IF(M{summary_row + 1}="","未提供图片重量",IF(M{summary_row + 4}<=0.01,"一致",IF(M{summary_row + 4}<=1,"差异≤1KG","有差异")))',
        summary_text,
        conclusion,
    )

    note_row = max(13, 4 + len(matches) + 2)
    worksheet.set_row(note_row, 38)
    worksheet.merge_range(
        note_row,
        0,
        note_row,
        8,
        "计算：单件重量（T）＝基础表总重量（T）÷基础件数；本次拆分重量（KG）＝单件重量（T）×图片拆分数量×1000。图片钢板重量仅用于复核，不用于反推或分摊。",
        note,
    )
    source_paths = sorted({match.source.source_path for match in matches})
    source_books = sorted({match.source.source_workbook for match in matches})
    source_text = (
        f"基础数据来源：{'；'.join(source_paths)}（优先通道：{source_paths[0].split('/')[1] if '/' in source_paths[0] else source_paths[0]}）；"
        f"汇总表：{'；'.join(source_books)}；原文件名：{image.original_filename}；匹配状态：唯一匹配"
    )
    worksheet.set_row(note_row + 1, 28)
    worksheet.merge_range(note_row + 1, 0, note_row + 1, 8, source_text, source_note)
    workbook.close()


def validate_result(path: Path, expected_rows: int) -> None:
    formula_wb = load_workbook(path, data_only=False)
    value_wb = load_workbook(path, data_only=True)
    if formula_wb.sheetnames != ["拆图明细"]:
        raise ValueError(f"工作表异常：{formula_wb.sheetnames}")
    formula_ws = formula_wb["拆图明细"]
    value_ws = value_wb["拆图明细"]
    if formula_ws["F2"].value != expected_rows:
        raise ValueError("Excel 零件条数与识别结果不一致")
    formulas = [cell.value for row in formula_ws.iter_rows() for cell in row if isinstance(cell.value, str) and cell.value.startswith("=")]
    if not formulas:
        raise ValueError("Excel 未保留公式")
    for row in formula_ws.iter_rows():
        for cell in row:
            value = cell.value
            if isinstance(value, str) and any(error in value.upper() for error in FORMULA_ERRORS):
                raise ValueError(f"Excel 公式错误：{cell.coordinate}={value}")
    for row in value_ws.iter_rows():
        for cell in row:
            value = cell.value
            if isinstance(value, str) and value.upper() in FORMULA_ERRORS:
                raise ValueError(f"Excel 缓存值错误：{cell.coordinate}={value}")

