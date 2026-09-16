from __future__ import annotations

import json
import tempfile
from pathlib import Path

import fitz
from openpyxl import load_workbook

from config import Settings
from drive_client import DriveClient
from excel_writer import validate_result, write_result
from image_parser import parse_image
from matcher import match_with_priority
from pdf_renderer import extract_pdf_text_pages
from pipeline import PDF_MIME_TYPE, _load_channel, _parse_pdf_by_page


SAMPLES = [
    "完成_废183.2 T70退0.pdf",
    "完成_废209.4 T80退0.pdf",
    "完成_#2327 T80退0.pdf",
    "完成_#2185 T80退0.pdf",
    "完成_#2334 T80退700X1650.pdf",
    "完成_#2326 T50退2230X4150.pdf",
    "完成_#2323 T150退400X1760.jpg",
    "完成_#2333 T60退0.jpg",
]


def _facts(path: Path, expected_rows: int) -> list[list[object]]:
    workbook = load_workbook(path, data_only=True)
    worksheet = workbook["拆图明细"]
    rows: list[list[object]] = []
    for row in range(5, 5 + expected_rows):
        rows.append([worksheet.cell(row, column).value for column in range(2, 10)])
    workbook.close()
    return rows


def _same_value(left: object, right: object) -> bool:
    if isinstance(left, (int, float)) and isinstance(right, (int, float)):
        return abs(float(left) - float(right)) <= 0.02
    return left == right


def _compare(generated: Path, existing: Path, expected_rows: int) -> list[str]:
    current = _facts(generated, expected_rows)
    historical = _facts(existing, expected_rows)
    problems: list[str] = []
    if len(current) != len(historical):
        return [f"行数不一致：当前={len(current)} 历史={len(historical)}"]
    columns = ["订单号", "图号", "厚度", "基础件数", "拆分数量", "基础总重量T", "单件重量T", "拆分重量KG"]
    for row_index, (new_row, old_row) in enumerate(zip(current, historical), start=1):
        for column_index, (new_value, old_value) in enumerate(zip(new_row, old_row)):
            if not _same_value(new_value, old_value):
                problems.append(
                    f"第{row_index}项 {columns[column_index]} 不一致：当前={new_value!r} 历史={old_value!r}"
                )
    return problems


def main() -> None:
    settings = Settings.from_env()
    drive = DriveClient()
    pending = drive.find_child_folder(settings.root_folder_id, settings.pending_folder_name)
    processing = drive.find_child_folder(settings.root_folder_id, settings.processing_folder_name)
    templates = drive.find_child_folder(settings.root_folder_id, settings.template_folder_name)
    results = drive.find_child_folder(settings.root_folder_id, settings.result_folder_name)

    pending_items = {item["name"]: item for item in drive.list_children(pending["id"])}
    result_items: dict[str, list[dict]] = {}
    for item in drive.list_children(results["id"]):
        result_items.setdefault(item["name"], []).append(item)

    missing = [name for name in SAMPLES if name not in pending_items]
    if missing:
        raise RuntimeError(f"历史源文件缺失：{missing}")

    report: list[dict] = []
    failures: list[str] = []
    with tempfile.TemporaryDirectory(prefix="historical-regression-") as temp_name:
        root = Path(temp_name)
        primary_parts = _load_channel(drive, processing["id"], settings.processing_folder_name, root / "sources")
        fallback_parts = _load_channel(drive, templates["id"], settings.template_folder_name, root / "sources")

        for source_name in SAMPLES:
            item = pending_items[source_name]
            original_name = source_name[len(settings.completed_prefix):] if source_name.startswith(settings.completed_prefix) else source_name
            suffix = Path(source_name).suffix.lower()
            local_input = root / "inputs" / source_name
            drive.download(item["id"], local_input)
            page_count = 1
            ocr_pages: list[int] = []
            route = "image-ocr"

            if item.get("mimeType") == PDF_MIME_TYPE or suffix == ".pdf":
                with fitz.open(local_input) as document:
                    page_count = document.page_count
                page_texts = extract_pdf_text_pages(local_input)
                image, ocr_pages = _parse_pdf_by_page(
                    original_name,
                    local_input,
                    page_texts,
                    root / "rendered" / item["id"],
                )
                route = "pdf-native" if not ocr_pages else "pdf-pagewise-ocr"
            else:
                image = parse_image(local_input, original_filename=original_name)

            matches = match_with_priority(image, primary_parts, fallback_parts)
            generated = root / "generated" / f"{image.main_name}_完成.xlsx"
            write_result(generated, image, matches)
            validate_result(generated, expected_rows=len(matches))

            result_name = generated.name
            existing_matches = result_items.get(result_name, [])
            if len(existing_matches) != 1:
                failures.append(f"{source_name}: 历史结果 {result_name} 应唯一，实际 {len(existing_matches)} 个")
                continue
            existing = root / "existing" / result_name
            drive.download(existing_matches[0]["id"], existing)
            differences = _compare(generated, existing, len(matches))
            if differences:
                failures.append(f"{source_name}: " + "；".join(differences[:8]))

            report.append(
                {
                    "source": source_name,
                    "type": suffix.lstrip("."),
                    "pages": page_count,
                    "route": route,
                    "ocr_pages": ocr_pages,
                    "parts": len(image.parts),
                    "program": image.program_no,
                    "marked_weight_kg": image.marked_weight_kg,
                    "output": result_name,
                    "historical_compare": "PASS" if not differences else "FAIL",
                    "differences": differences,
                }
            )

    print("HISTORICAL_REGRESSION_REPORT=" + json.dumps(report, ensure_ascii=False))
    print(f"HISTORICAL_REGRESSION_SUMMARY=total:{len(report)},pass:{sum(1 for item in report if item['historical_compare']=='PASS')},fail:{len(failures)}")
    if failures:
        print("HISTORICAL_REGRESSION_FAILURES=" + json.dumps(failures, ensure_ascii=False))
        raise SystemExit(1)


if __name__ == "__main__":
    main()
