from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from openpyxl import Workbook, load_workbook

from excel_writer import validate_result, write_result
from image_parser import _parse_parts, _parse_program, _parse_weight, main_name_from_filename, parse_image
from matcher import MatchError, match_image, match_one
from models import ImageData, ImagePart, SourcePart
from source_reader import SourceReadError, read_summary_workbook


OCR_2323 = """
1 THP11-6000L-0805 1001-01-08 T150 2) W x 2
2 THP11-6000L-0805 1001-01-17 T150 2) W x 2
重量 3041.02kg
日期 08:38:50 2026/9/5 N163
"""


class CoreTests(unittest.TestCase):
    def test_2323_ocr_text(self) -> None:
        parts = _parse_parts(OCR_2323)
        self.assertEqual([part.drawing_no for part in parts], ["1001-01-08", "1001-01-17"])
        self.assertEqual([part.base_quantity for part in parts], [2, 2])
        self.assertEqual([part.split_quantity for part in parts], [2, 2])
        self.assertEqual(_parse_weight(OCR_2323), 3041.02)
        self.assertEqual(_parse_program(OCR_2323), "N163")

    def test_common_ocr_confusions(self) -> None:
        self.assertEqual(_parse_weight("重量 14307.%ke"), 14307.96)
        self.assertEqual(_parse_program("代码文件名 NI84"), "N184")

    def test_chinese_named_part_without_order_number(self) -> None:
        text = """
1 YT71S-2000WA-0711 1007-01-02 T40 4J P x 2
2 吊 耳 T40 100J P x 33
3 YT71S-2500Z-0715 1004-01-16 T40 4J P x 4
4 YT71S-2000WA-0711 1004-01-14.2 T40 4J P x 1
"""
        parts = _parse_parts(text)
        self.assertEqual([part.index for part in parts], [1, 2, 3, 4])
        self.assertEqual(parts[1].order_no, "")
        self.assertEqual(parts[1].drawing_no, "吊耳")
        self.assertEqual(parts[1].base_quantity, 100)
        self.assertEqual(parts[1].split_quantity, 33)

    def test_chinese_named_part_requires_unique_source_match(self) -> None:
        part = ImagePart(2, "", "吊耳", 40, 100, "P", 33)
        source = SourcePart("YT71S-2000WA-0711", "吊耳", 40, 100, "P", 1234, "路径", "汇总表.xlsx")
        self.assertEqual(match_one(part, [source]), source)
        duplicate = SourcePart("OTHER-ORDER-0001", "吊耳", 40, 100, "P", 1234, "其他路径", "汇总表.xlsx")
        with self.assertRaises(MatchError):
            match_one(part, [source, duplicate])

    def test_unreadable_name_falls_back_to_unique_physical_fields(self) -> None:
        parts = _parse_parts("1 ib A T40 100J P x 33")
        self.assertEqual(parts[0], ImagePart(1, "", "", 40, 100, "P", 33))
        source = SourcePart("YT71S-2000WA-0711", "吊耳", 40, 100, "P", 1234, "路径", "汇总表.xlsx")
        self.assertEqual(match_one(parts[0], [source]), source)
        duplicate = SourcePart("OTHER-ORDER-0001", "其他吊耳", 40, 100, "P", 900, "其他路径", "汇总表.xlsx")
        with self.assertRaises(MatchError):
            match_one(parts[0], [source, duplicate])

    def test_short_order_and_drawing_codes(self) -> None:
        text = """
1 05TD1 05TD1-01 T60 1J P x1
2 05TD1 05TD2-01 T60 1J P x1
3 05TD1 05TD2-03 T60 1J W x1
"""
        parts = _parse_parts(text)
        self.assertEqual([part.order_no for part in parts], ["05TD1"] * 3)
        self.assertEqual([part.drawing_no for part in parts], ["05TD1-01", "05TD2-01", "05TD2-03"])

    def test_historical_images_when_local_fixtures_exist(self) -> None:
        fixtures = Path(__file__).parent / "fixtures"
        expected = {
            "2323": ("N163", 3041.02, 2),
            "2324": ("N167", 7480.08, 2),
            "2325": ("N167", 5193.93, 1),
            "2326": ("N172", 2608.06, 1),
            "2331": ("N184", 14307.96, 2),
            "2332": ("N186", 13677.45, 2),
            "2333": ("N196", 2659.94, 1),
        }
        available = [key for key in expected if (fixtures / f"{key}.jpg").exists()]
        if not available:
            self.skipTest("生产历史图片未提交到仓库")
        for key in available:
            with self.subTest(image=key):
                parsed = parse_image(fixtures / f"{key}.jpg", f"#{key} 测试.jpg")
                program, weight, part_count = expected[key]
                self.assertEqual(parsed.program_no, program)
                self.assertAlmostEqual(parsed.marked_weight_kg, weight)
                self.assertEqual(len(parsed.parts), part_count)

    def test_old_output_filename(self) -> None:
        self.assertEqual(main_name_from_filename("#2323 T150退400X1760.jpg"), "#2323")
        self.assertEqual(main_name_from_filename("完成_废4.5 T30退0.jpg"), "废4.5")

    def test_2323_matching_weight_and_excel(self) -> None:
        with tempfile.TemporaryDirectory() as name:
            directory = Path(name)
            summary = directory / "summary.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["订单号", "图号", "厚度", "件数", "长(mm)", "宽(mm)", "切割长度(mm)", "净面积(m²)", "坡口", "总净重(kg)", "核对备注"])
            for drawing in ["1001-01-08", "1001-01-17"]:
                sheet.append(["174.26-08-05  THP11-6000L-0805", drawing, 150, 2, 900, 850, 4778.6, 0.634899579631376, "W", 1495.18851003189, "未发现异常差异"])
            workbook.save(summary)
            sources = read_summary_workbook(summary, "王振海/正在加工/174")
            image = ImageData(
                original_filename="#2323 T150退400X1760.jpg",
                main_name="#2323",
                program_no="N163",
                marked_weight_kg=3041.02,
                parts=_parse_parts(OCR_2323),
                ocr_text=OCR_2323,
            )
            matches = match_image(image, sources)
            self.assertAlmostEqual(sum(item.split_weight_kg for item in matches), 2990.37702006378)
            result = directory / "#2323_完成.xlsx"
            write_result(result, image, matches)
            validate_result(result, 2)
            formula_sheet = load_workbook(result, data_only=False).active
            value_sheet = load_workbook(result, data_only=True).active
            self.assertEqual(formula_sheet["H5"].value, "=G5/E5")
            self.assertEqual(formula_sheet["I5"].value, "=H5*F5*1000")
            self.assertAlmostEqual(value_sheet["M9"].value, 2990.37702006378)
            self.assertEqual(value_sheet["M12"].value, "有差异")

    def test_strict_match_rejects_missing_source(self) -> None:
        image = ImageData(
            "x.jpg",
            "x",
            "N1",
            1.0,
            (ImagePart(1, "THP11-6000L-0805", "1001-01-08", 150, 2, "P", 2),),
            "",
        )
        with self.assertRaises(MatchError):
            match_image(image, [])

    def test_special_order_uses_unique_drawing_tail(self) -> None:
        image = ImageData(
            "special.jpg",
            "special",
            "N1",
            1.0,
            (ImagePart(1, "26511255302-0814", "AA-01-08", 50, 2, "P", 1),),
            "",
        )
        source = SourcePart(
            "26511255302-0814", "LONG-PREFIX-01-08", 50, 2, "P", 1000,
            "王振海/拆图模版/179", "汇总表.xlsx",
        )
        self.assertEqual(match_image(image, [source])[0].source, source)

    def test_xlsm_summary_weight_is_tonnes(self) -> None:
        with tempfile.TemporaryDirectory() as name:
            path = Path(name) / "模板.xlsm"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["订单号", "图号", "厚度", "件数", "长", "宽", "切割长度", "面积", "坡口", "重量"])
            sheet.append(["JYT27-1600E-0812", "1001-01-45A", "50", "2", "1", "1", "1", "1", "P", 3.536])
            workbook.save(path)
            parts = read_summary_workbook(path, "王振海/正在加工/177")
            self.assertEqual(parts[0].total_weight_kg, 3536.0)

    def test_unlabelled_structured_bevel_column(self) -> None:
        with tempfile.TemporaryDirectory() as name:
            path = Path(name) / "无坡口标题模板.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["订单号", "图号", "厚度", "件数", "长", "宽", "切割长度", "面积", "", "重量"])
            sheet.append(["THP10-3150F-0910", "19Y01-01", 100, 8, 160, 160, 502.65, 0.02011, "P", 0.126])
            sheet.append(["THP10-3150F-0910", "19Y01-02", 50, 4, 220, 1080, 2882.74, 0.23442, "W", 0.368])
            workbook.save(path)
            parts = read_summary_workbook(path, "王振海/正在加工/THP10")
            self.assertEqual([part.bevel for part in parts], ["P", "W"])
            self.assertEqual([part.total_weight_kg for part in parts], [126.0, 368.0])

    def test_empty_sheet_before_summary_is_skipped(self) -> None:
        with tempfile.TemporaryDirectory() as name:
            path = Path(name) / "含空表模板.xlsx"
            workbook = Workbook()
            workbook.active.title = "空表"
            sheet = workbook.create_sheet("汇总")
            sheet.append(["订单号", "图号", "厚度", "件数", "坡口", "重量"])
            sheet.append(["THP10-3150F-0910", "19Y01-01", 100, 8, "P", 0.126])
            workbook.save(path)
            parts = read_summary_workbook(path, "王振海/正在加工/THP10")
            self.assertEqual(len(parts), 1)
            self.assertEqual(parts[0].drawing_no, "19Y01-01")

    def test_missing_cached_weight_is_rejected(self) -> None:
        with tempfile.TemporaryDirectory() as name:
            path = Path(name) / "汇总表.xlsx"
            workbook = Workbook()
            sheet = workbook.active
            sheet.append(["订单号", "图号", "厚度", "件数", "坡口", "总净重(kg)"])
            sheet.append(["THP11-6000L-0805", "1001-01-08", 150, 2, "W", "=1+1"])
            workbook.save(path)
            with self.assertRaises(SourceReadError):
                read_summary_workbook(path, "王振海/正在加工/174")


if __name__ == "__main__":
    unittest.main()
