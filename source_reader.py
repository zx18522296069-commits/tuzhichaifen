from __future__ import annotations

import re
from pathlib import Path

from openpyxl import load_workbook

from models import SourcePart


class SourceReadError(RuntimeError):
    pass


def _normal(value: object) -> str:
    return (
        str(value or "")
        .replace("\u200b", "")
        .replace("\ufeff", "")
        .strip()
        .upper()
        .replace("—", "-")
        .replace("–", "-")
    )


def _order_number(value: object) -> str:
    normalized = _normal(value)
    candidates = re.findall(r"(?:[A-Z0-9]+(?:-[A-Z0-9]+){2,}|\d{8,}-\d{4})", normalized)
    # 汇总表里这一列的业务含义是“型号/订单”。有些订单只写 05TD1
    # 一类短型号，不能因为它不含两个连字符就丢弃这一整行。
    if candidates:
        return candidates[-1]
    if normalized:
        return normalized
    raise SourceReadError(f"无法从型号/订单字段读取内容：{value!r}")


def _header_key(value: object) -> str:
    """Return a semantic field key without relying on a fixed workbook layout."""
    text = re.sub(r"[\s_\-()（）\[\]【】]", "", _normal(value))
    if text in {"订单号", "订单编号", "型号", "产品型号", "机型", "型号订单号"}:
        return "order"
    if text in {"图号", "图纸号", "零件号", "零件图号", "零件编号"}:
        return "drawing"
    if text in {"厚度", "板厚", "厚度MM", "板厚MM"}:
        return "thickness"
    if text in {"件数", "数量", "件数件", "数量件"}:
        return "quantity"
    if text in {"坡口", "坡口形式"}:
        return "bevel"
    if text in {"总净重KG", "总重量KG", "净重KG", "重量KG"}:
        return "weight_kg"
    if text in {"总净重T", "总重量T", "净重T", "重量T", "重量"}:
        return "weight_t"
    return ""


def read_summary_workbook(path: Path, source_path: str) -> list[SourcePart]:
    keep_vba = path.suffix.lower() == ".xlsm"
    formula_wb = load_workbook(path, data_only=False, read_only=True, keep_vba=keep_vba)
    value_wb = load_workbook(path, data_only=True, read_only=True, keep_vba=keep_vba)

    header_row = None
    headers: dict[str, int] = {}
    value_ws = None
    for worksheet in value_wb.worksheets:
        # Some xlsm exports omit the worksheet dimension.  In read-only mode
        # openpyxl then reports max_row=None, even though row 1 contains the
        # summary header.  Always stream the first 20 physical rows instead.
        for row in worksheet.iter_rows(min_row=1, max_row=20):
            semantic_headers = {
                key: index + 1
                for index, cell in enumerate(row)
                if (key := _header_key(cell.value))
            }
            if {"order", "drawing", "thickness", "quantity"}.issubset(semantic_headers):
                value_ws = worksheet
                header_row = row[0].row
                headers = semantic_headers
                break
        if value_ws is not None:
            break
    if value_ws is None or not header_row:
        raise SourceReadError(f"{path.name} 未找到汇总表表头")
    formula_ws = formula_wb[value_ws.title]

    def column(name: str) -> int:
        if name in headers:
            return headers[name]
        labels = {
            "order": "型号/订单号",
            "drawing": "图号",
            "thickness": "厚度",
            "quantity": "件数",
        }
        raise SourceReadError(f"{path.name} 缺少字段：{labels.get(name, name)}")

    order_col = column("order")
    drawing_col = column("drawing")
    thickness_col = column("thickness")
    quantity_col = column("quantity")
    weight_col = headers.get("weight_kg") or headers.get("weight_t")
    weight_multiplier = 1.0 if "weight_kg" in headers else 1000.0
    if weight_col is None or weight_multiplier is None:
        raise SourceReadError(f"{path.name} 缺少总重量字段（kg、t 或模板重量）")

    bevel_col = headers.get("bevel")
    if bevel_col is None:
        # Some production summary sheets keep the bevel values in the normal
        # detail column but leave that column's heading blank.  Only accept a
        # single column between quantity and weight whose populated detail
        # values are unambiguous bevel codes; otherwise stop instead of
        # guessing.
        candidates: list[int] = []
        for candidate in range(quantity_col + 1, weight_col):
            values: list[str] = []
            for row in value_ws.iter_rows(min_row=header_row + 1):
                if not _normal(row[drawing_col - 1].value):
                    continue
                value = _normal(row[candidate - 1].value)
                if value:
                    values.append(value)
            if values and all(re.fullmatch(r"[A-Z][A-Z0-9]{0,3}", value) for value in values):
                candidates.append(candidate)
        if len(candidates) != 1:
            raise SourceReadError(f"{path.name} 缺少字段：坡口")
        bevel_col = candidates[0]
    result: list[SourcePart] = []
    for row_no, row in enumerate(value_ws.iter_rows(min_row=header_row + 1), start=header_row + 1):
        drawing = _normal(row[drawing_col - 1].value)
        if not drawing:
            continue
        try:
            order = _order_number(row[order_col - 1].value)
            thickness = float(row[thickness_col - 1].value)
            quantity = int(row[quantity_col - 1].value)
        except (TypeError, ValueError, SourceReadError):
            continue
        bevel = _normal(row[bevel_col - 1].value)
        cached_weight = row[weight_col - 1].value
        if isinstance(cached_weight, (int, float)):
            total_weight_kg = float(cached_weight) * weight_multiplier
        else:
            formula = formula_ws.cell(row_no, weight_col).value
            raise SourceReadError(f"{path.name} 第 {row_no} 行总重量没有可读取的已确认值：{formula!r}")
        result.append(
            SourcePart(
                order_no=order,
                drawing_no=drawing,
                thickness=thickness,
                base_quantity=quantity,
                bevel=bevel,
                total_weight_kg=total_weight_kg,
                source_path=source_path,
                source_workbook=path.name,
            )
        )
    return result
